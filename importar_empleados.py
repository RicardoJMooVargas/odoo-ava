#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de importación masiva de empleados y documentos para Odoo 19.
Módulo: expediente_empleados

Uso:
  python importar_empleados.py --user tu_usuario@email.com
  python importar_empleados.py --user tu_usuario@email.com --dry-run
"""

import os
import re
import sys
import getpass
import argparse
import base64
import datetime
import unicodedata
import xmlrpc.client
import openpyxl


def normalizar_texto(texto):
    """Normaliza texto removiendo acentos, guiones bajos y espacios extras."""
    if not texto:
        return ""
    texto = str(texto).replace("_", " ").strip()
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("utf-8")
    return re.sub(r"\s+", " ", texto.lower())


def mapear_estatus(estatus_raw):
    """Mapea el estatus de texto al valor técnico del campo estatus_expediente."""
    if not estatus_raw:
        return "borrador"
    val = str(estatus_raw).strip().lower()
    if "activo" in val or "vigente" in val:
        return "vigente"
    elif "inactivo" in val or "baja" in val:
        return "baja"
    return "borrador"


def limpiar_nss(nss_raw):
    """Limpia el número de seguridad social (evita decimales como .0)."""
    if nss_raw is None:
        return ""
    if isinstance(nss_raw, (int, float)):
        return str(int(nss_raw))
    val = str(nss_raw).strip()
    if val.endswith(".0"):
        val = val[:-2]
    return val


def cargar_datos_excel(excel_path):
    """Lee el archivo db_empleados.xlsx y retorna una lista de diccionarios con los empleados."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"No se encontró el archivo Excel: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "empleados" not in wb.sheetnames:
        raise ValueError("El archivo Excel no contiene una hoja llamada 'empleados'")

    ws = wb["empleados"]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    if not rows:
        raise ValueError("La hoja 'empleados' está vacía.")

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    empleados = []

    # Índices conocidos
    col_rfc = headers.index("rfc") if "rfc" in headers else 18

    for row_idx, r in enumerate(rows[1:], start=2):
        colaborador = str(r[2]).strip() if r[2] else ""
        if not colaborador:
            continue

        fecha_ing = r[1]
        if isinstance(fecha_ing, (datetime.datetime, datetime.date)):
            fecha_ing_str = fecha_ing.strftime("%Y-%m-%d")
        elif fecha_ing:
            fecha_ing_str = str(fecha_ing).strip()[:10]
        else:
            fecha_ing_str = False

        # RFC puede venir en la columna 'rfc' o en la columna adicional (21)
        rfc_val = r[col_rfc] if col_rfc < len(r) else None
        if not rfc_val and len(r) > 21 and r[21]:
            rfc_val = r[21]
        rfc_str = str(rfc_val).strip() if rfc_val else ""

        emp = {
            "excel_row": row_idx,
            "id_plantilla": int(r[0]) if r[0] else None,
            "name": colaborador,
            "norm_name": normalizar_texto(colaborador),
            "fecha_ingreso": fecha_ing_str,
            "puesto_expediente": str(r[3]).strip() if r[3] else "",
            "job_title": str(r[3]).strip() if r[3] else "",
            "es_chofer": bool(r[4]) if r[4] is not None else False,
            "observaciones_expediente": str(r[5]).strip() if r[5] else "",
            "division": str(r[6]).strip() if r[6] else "",
            "estatus_expediente": mapear_estatus(r[7]),
            "no_seguro_social": limpiar_nss(r[8]),
            "rfc": rfc_str,
            "documentos": {}
        }
        empleados.append(emp)

    return empleados


def clasificar_archivos(docs_folder, empleados):
    """
    Escanea la carpeta subida_documents y asocia cada archivo al empleado y campo Many2many correspondiente.
    """
    if not os.path.isdir(docs_folder):
        print(f"[ADVERTENCIA] No se encontró la carpeta de documentos: {docs_folder}")
        return

    # Mapeo de prefijo del archivo -> campo Many2many en hr.employee
    field_prefix_map = {
        "ACTA_NAC_": "act_naci_ids",
        "COMP_DOMIC_": "comp_domic_ids",
        "CONST_FISC_": "const_sit_fiscal_ids",
        "CONTRATO_": "contrato_ids",
        "CURP_": "curp_ids",
        "DESC_PUESTO_": "descript_puesto_ids",
        "INE_": "ine_ids",
        "SOLICITUD_": "solicitud_empl_ids",
        "CV_": "chatter_cv",  # CV se adjunta al chatter ya que no tiene campo propio
    }

    archivos = sorted(os.listdir(docs_folder))
    asignados = 0
    duplicados_saltados = 0

    for f in archivos:
        full_path = os.path.join(docs_folder, f)
        if not os.path.isfile(full_path):
            continue

        # Evitar copias duplicadas de descarga tipo archivo(1).pdf si existe el archivo base
        if re.search(r"\(\d+\)\.pdf$", f, re.IGNORECASE):
            base_f = re.sub(r"\(\d+\)\.pdf$", ".pdf", f, flags=re.IGNORECASE)
            if os.path.exists(os.path.join(docs_folder, base_f)):
                duplicados_saltados += 1
                continue

        # Identificar prefijo
        matched_prefix = None
        for prefix in field_prefix_map:
            if f.upper().startswith(prefix):
                matched_prefix = prefix
                break

        if not matched_prefix:
            continue

        # Extraer nombre del empleado en el archivo
        nombre_en_archivo = f[len(matched_prefix):]
        nombre_en_archivo = os.path.splitext(nombre_en_archivo)[0]
        norm_archivo = normalizar_texto(nombre_en_archivo)

        # Buscar empleado por coincidencia normalizada
        empleado_encontrado = None
        for emp in empleados:
            if (emp["norm_name"] == norm_archivo or
                emp["norm_name"] in norm_archivo or
                norm_archivo in emp["norm_name"]):
                empleado_encontrado = emp
                break

        if empleado_encontrado:
            m2m_field = field_prefix_map[matched_prefix]
            empleado_encontrado["documentos"].setdefault(m2m_field, []).append({
                "filename": f,
                "path": full_path,
                "prefix": matched_prefix
            })
            asignados += 1
        else:
            print(f"  [AVISO] No se encontró empleado para archivo: {f}")

    print(f"-> Archivos escaneados: {len(archivos)}")
    print(f"-> Archivos asociados con éxito: {asignados}")
    if duplicados_saltados:
        print(f"-> Archivos duplicados saltados: {duplicados_saltados}")


def conectar_odoo(url, db, user, password):
    """Establece conexión XML-RPC con Odoo y autentica."""
    common_url = f"{url.rstrip('/')}/xmlrpc/2/common"
    object_url = f"{url.rstrip('/')}/xmlrpc/2/object"

    common = xmlrpc.client.ServerProxy(common_url)
    uid = common.authenticate(db, user, password, {})
    if not uid:
        raise PermissionError(f"Error de autenticación en Odoo para el usuario '{user}' en DB '{db}'")

    models = xmlrpc.client.ServerProxy(object_url)
    return uid, models


def main():
    parser = argparse.ArgumentParser(description="Importador masivo de empleados y documentos a Odoo 19.")
    parser.add_argument("--url", default="https://odoo.vyaainfor.systems", help="URL de la instancia Odoo")
    parser.add_argument("--db", default="odoo", help="Nombre de la base de datos de Odoo")
    parser.add_argument("--user", help="Usuario o Email de Odoo")
    parser.add_argument("--password", help="Contraseña o API Key de Odoo")
    parser.add_argument("--excel", default="db_empleados.xlsx", help="Ruta al archivo db_empleados.xlsx")
    parser.add_argument("--docs", default="subida_documents", help="Ruta a la carpeta subida_documents")
    parser.add_argument("--dry-run", action="store_true", help="Simula el proceso sin escribir en Odoo")
    parser.add_argument("--skip-docs", action="store_true", help="Solo importa/actualiza datos de empleados sin adjuntos")

    args = parser.parse_args()

    print("=" * 70)
    print("  IMPORTADOR MASIVO DE EMPLEADOS Y DOCUMENTOS PARA ODOO 19")
    print("=" * 70)

    # 1. Cargar Excel
    print(f"\n[1/4] Leyendo archivo Excel: {args.excel}")
    empleados = cargar_datos_excel(args.excel)
    print(f"-> Total colaboradores encontrados en Excel: {len(empleados)}")

    # 2. Clasificar archivos locales
    if not args.skip_docs:
        print(f"\n[2/4] Escaneando carpeta de documentos: {args.docs}")
        clasificar_archivos(args.docs, empleados)
    else:
        print("\n[2/4] Omitiendo documentos por bandera --skip-docs")

    # 3. Credenciales de Odoo
    print(f"\n[3/4] Conectando a Odoo ({args.url})...")
    user = args.user or input("Usuario / Email de Odoo: ").strip()
    password = args.password or getpass.getpass("Contraseña / API Key de Odoo: ")

    if not args.dry_run:
        try:
            uid, models = conectar_odoo(args.url, args.db, user, password)
            print(f"-> Autenticado con éxito en Odoo (UID: {uid})")
        except Exception as e:
            print(f"\n[ERROR] No se pudo conectar a Odoo: {e}")
            sys.exit(1)
    else:
        print("-> [MODO SIMULACIÓN --dry-run ACTIVO: No se realizarán cambios en Odoo]")
        uid, models = None, None

    # 4. Proceso de subida
    print(f"\n[4/4] Procesando empleados en Odoo...")
    creados = 0
    actualizados = 0
    adjuntos_subidos = 0

    for idx, emp in enumerate(empleados, start=1):
        nombre = emp["name"]
        print(f"\n[{idx}/{len(empleados)}] {nombre}")

        # Campos del empleado para Odoo
        vals_empleado = {
            "name": nombre,
            "fecha_ingreso": emp["fecha_ingreso"] or False,
            "puesto_expediente": emp["puesto_expediente"],
            "job_title": emp["job_title"],
            "es_chofer": emp["es_chofer"],
            "observaciones_expediente": emp["observaciones_expediente"],
            "division": emp["division"],
            "estatus_expediente": emp["estatus_expediente"],
            "no_seguro_social": emp["no_seguro_social"],
        }
        if emp["rfc"]:
            vals_empleado["rfc"] = emp["rfc"]

        emp_id = None
        if not args.dry_run:
            # Buscar si el empleado ya existe en Odoo por nombre
            existentes = models.execute_kw(
                args.db, uid, password,
                "hr.employee", "search",
                [[["name", "=ilike", nombre]]]
            )

            if existentes:
                emp_id = existentes[0]
                models.execute_kw(
                    args.db, uid, password,
                    "hr.employee", "write",
                    [[emp_id], vals_empleado]
                )
                actualizados += 1
                print(f"  -> Actualizado empleado existente (ID Odoo: {emp_id})")
            else:
                emp_id = models.execute_kw(
                    args.db, uid, password,
                    "hr.employee", "create",
                    [vals_empleado]
                )
                creados += 1
                print(f"  -> Creado nuevo empleado (ID Odoo: {emp_id})")
        else:
            print(f"  [SIMULACIÓN] Se crearía/actualizaría empleado: {vals_empleado['name']} ({vals_empleado['puesto_expediente']})")
            emp_id = 9999

        # Procesar documentos del empleado
        if not args.skip_docs and emp["documentos"]:
            for campo_m2m, lista_archivos in emp["documentos"].items():
                for doc in lista_archivos:
                    fname = doc["filename"]
                    fpath = doc["path"]

                    if args.dry_run:
                        print(f"  [SIMULACIÓN] Adjuntaría archivo '{fname}' al campo '{campo_m2m}'")
                        adjuntos_subidos += 1
                        continue

                    try:
                        # Verificar si el adjunto ya existe para evitar duplicados
                        att_existentes = models.execute_kw(
                            args.db, uid, password,
                            "ir.attachment", "search",
                            [[
                                ["res_model", "=", "hr.employee"],
                                ["res_id", "=", emp_id],
                                ["name", "=", fname]
                            ]]
                        )

                        if att_existentes:
                            att_id = att_existentes[0]
                            # Asegurar vinculación Many2many si corresponde
                            if campo_m2m != "chatter_cv":
                                models.execute_kw(
                                    args.db, uid, password,
                                    "hr.employee", "write",
                                    [[emp_id], {campo_m2m: [(4, att_id)]}]
                                )
                            print(f"    - Adjunto existente verificado: {fname}")
                            continue

                        # Leer archivo y codificar en Base64
                        with open(fpath, "rb") as arch:
                            contenido_b64 = base64.b64encode(arch.read()).decode("utf-8")

                        # Crear ir.attachment
                        vals_att = {
                            "name": fname,
                            "datas": contenido_b64,
                            "res_model": "hr.employee",
                            "res_id": emp_id,
                            "mimetype": "application/pdf"
                        }
                        att_id = models.execute_kw(
                            args.db, uid, password,
                            "ir.attachment", "create",
                            [vals_att]
                        )

                        # Si no es CV (que queda como adjunto general), vincular al campo Many2many
                        if campo_m2m != "chatter_cv":
                            models.execute_kw(
                                args.db, uid, password,
                                "hr.employee", "write",
                                [[emp_id], {campo_m2m: [(4, att_id)]}]
                            )

                        adjuntos_subidos += 1
                        print(f"    + Subido y vinculado: {fname} -> {campo_m2m}")

                    except Exception as err_doc:
                        print(f"    [ERROR] Falló subida de {fname}: {err_doc}")

    # Resumen final
    print("\n" + "=" * 70)
    print("  RESUMEN DE IMPORTACIÓN")
    print("=" * 70)
    if not args.dry_run:
        print(f"Empleados creados:      {creados}")
        print(f"Empleados actualizados:  {actualizados}")
        print(f"Documentos subidos:      {adjuntos_subidos}")
        print("\n¡Proceso finalizado exitosamente!")
    else:
        print(f"Empleados simulados:     {len(empleados)}")
        print(f"Documentos simulados:    {adjuntos_subidos}")
        print("\n[SIMULACIÓN FINALIZADA: Todo listo para ejecutar en modo real]")


if __name__ == "__main__":
    main()
